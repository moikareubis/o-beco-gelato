# -*- coding: utf-8 -*-
"""CMV por produto — planilha viva. Tudo que é amarelo o Gabriel edita; o resto é fórmula."""
import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = json.load(open("_dossie.json", encoding="utf-8"))
HOJE = datetime.date(2026, 10, 21)

MARROM="6B4A2F"; AREIA="F3EDE5"; CREME="FBF6EF"; CINZA="6E655C"
AMARELO="FFF4CC"; VERDE="EDF5F0"; VERM="FBEDEA"
NEG=Font(bold=True, color="1F1A17", size=10)
CAB=Font(bold=True, color="FFFFFF", size=10)
TIT=Font(bold=True, color=MARROM, size=13)
PEQ=Font(size=8, color=CINZA, italic=True)
FCAB=PatternFill("solid", fgColor=MARROM)
FEDIT=PatternFill("solid", fgColor=AMARELO)
FTOT=PatternFill("solid", fgColor=CREME)
FOK=PatternFill("solid", fgColor=VERDE)
BF=Border(bottom=Side(style="thin", color="D8CEC1"))
BOX=Border(*[Side(style="thin", color="B8AC9C")]*4)

wb = Workbook()
def cab(ws, l, tits, larg=None, altura=26):
    for i,t in enumerate(tits, start=1):
        c=ws.cell(row=l, column=i, value=t); c.font=CAB; c.fill=FCAB
        c.alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)
    if larg:
        for i,w_ in enumerate(larg, start=1):
            ws.column_dimensions[get_column_letter(i)].width=w_
    ws.row_dimensions[l].height=altura

SABS = D["sabores"]
NI = len(D["ingredientes"])

# ═════════ 1. PARÂMETROS — tudo que ele edita fica aqui
ws = wb.active; ws.title="PARÂMETROS"
ws.column_dimensions["A"].width=34; ws.column_dimensions["B"].width=14
ws.column_dimensions["C"].width=12; ws.column_dimensions["D"].width=62
ws["A1"]="CMV por produto — O Beco Gelato"; ws["A1"].font=TIT
ws["A2"]=f"Revisão de {HOJE.strftime('%d/%m/%Y')}. As células amarelas são suas — mude e tudo recalcula."
ws["A2"].font=PEQ

cab(ws, 4, ["Parâmetro","Valor","Unidade","O que é"])
params = [
 ("Gelato por pote", D["pote"], "g", "Medido em 16/10 com 35 potes. O Alfajor tem 135 g e é tratado à parte."),
 ("Perda de processo", 3, "%", "O que fica na cuba, na pá e na espátula. Medido a 3% em lote grande; em lote de 1,5 kg chega a 10%."),
 ("Imposto (Simples)", 4.5, "%", "Anexo II, 1ª faixa. CONFIRMAR com o contador."),
]
r=5
for n,v,u,o in params:
    ws.cell(row=r, column=1, value=n).font=NEG
    c=ws.cell(row=r, column=2, value=v); c.fill=FEDIT; c.border=BOX; c.number_format="0.0"
    ws.cell(row=r, column=3, value=u)
    c2=ws.cell(row=r, column=4, value=o); c2.alignment=Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height=28; r+=1
POTE_REF="PARÂMETROS!$B$5"; PERDA_REF="PARÂMETROS!$B$6"; IMP_REF="PARÂMETROS!$B$7"
r+=1

ws.cell(row=r, column=1, value="EMBALAGEM — por pote").font=NEG; r+=1
cab(ws, r, ["Item","Qtd por pote","R$ unitário","R$ por pote"]); r+=1
emb_ini=r
for e in D["embItens"]:
    ws.cell(row=r, column=1, value=e["nome"])
    c=ws.cell(row=r, column=2, value=e["porPote"]); c.fill=FEDIT; c.border=BOX
    c=ws.cell(row=r, column=3, value=e["preco"]); c.fill=FEDIT; c.border=BOX; c.number_format='"R$ "0.00'
    ws.cell(row=r, column=4, value=f"=B{r}*C{r}").number_format='"R$ "0.00'
    r+=1
emb_fim=r-1
ws.cell(row=r, column=1, value="TOTAL EMBALAGEM").font=NEG
c=ws.cell(row=r, column=4, value=f"=SUM(D{emb_ini}:D{emb_fim})")
c.font=NEG; c.number_format='"R$ "0.00'
for col in range(1,5): ws.cell(row=r, column=col).fill=FTOT
EMB_REF=f"PARÂMETROS!$D${r}"
r+=2

ws.cell(row=r, column=1, value="COMO A CONTA É FEITA").font=NEG; r+=1
for t in ["Calda por pote = gelato por pote ÷ (1 − perda)",
          "Custo de insumo = soma de (g de cada ingrediente no pote × R$/kg ÷ 1000)",
          "CMV = insumo + embalagem",
          "Imposto = preço de venda × alíquota",
          "Contribuição = preço − CMV − imposto",
          "Margem = contribuição ÷ preço"]:
    ws.cell(row=r, column=1, value=t); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4); r+=1

# ═════════ 2. PREÇOS DOS INSUMOS
ws = wb.create_sheet("PREÇOS")
ws["A1"]="Preço dos insumos"; ws["A1"].font=TIT
ws["A2"]="Coluna amarela: edite aqui quando o preço mudar. Todos os produtos recalculam."
ws["A2"].font=PEQ
cab(ws, 4, ["Insumo","R$ por kg","Observação"], [48,14,56])
r=5
for i in D["ingredientes"]:
    ws.cell(row=r, column=1, value=i["nome"])
    c=ws.cell(row=r, column=2, value=i["custoKg"]); c.fill=FEDIT; c.border=BOX
    c.number_format='"R$ "0.00'
    obs = "SEM PREÇO — a conta sai menor que o real" if not i["custoKg"] else ""
    c2=ws.cell(row=r, column=3, value=obs)
    if obs:
        for col in range(1,4): ws.cell(row=r, column=col).fill=PatternFill("solid", fgColor=VERM)
    for col in range(1,4): ws.cell(row=r, column=col).border=BF
    r+=1
ws.freeze_panes="A5"
PRECO_TAB=f"PREÇOS!$A$5:$B${r-1}"

# ═════════ 3. PREÇO DE VENDA
ws = wb.create_sheet("VENDA")
ws["A1"]="Preço de venda por produto"; ws["A1"].font=TIT
ws["A2"]="Preço praticado na rua, confirmado em 21/10. Edite para simular."
ws["A2"].font=PEQ
cab(ws, 4, ["Produto","Linha","R$ atacado","R$ varejo sugerido"], [26,12,14,18])
r=5
venda_lin={}
for s in SABS:
    ws.cell(row=r, column=1, value=s["nome"])
    ws.cell(row=r, column=2, value=s["linha"])
    c=ws.cell(row=r, column=3, value=s["preco"]); c.fill=FEDIT; c.border=BOX
    c.number_format='"R$ "0.00'
    vj = 29.90 if s["linha"]=="FRESH" else 32.90
    c=ws.cell(row=r, column=4, value=vj); c.fill=FEDIT; c.border=BOX
    c.number_format='"R$ "0.00'
    venda_lin[s["nome"]]=r
    r+=1

# ═════════ 4. UMA ABA POR PRODUTO
aba_total={}
for s in SABS:
    nome=s["nome"][:28]
    ws=wb.create_sheet("CMV "+nome[:24])
    ws["A1"]=f"CMV — {s['nome']}"; ws["A1"].font=TIT
    ws["A2"]=(f"linha {s['linha']} · {s['gelatoPote']:.0f} g de gelato por pote"
              + (f" + {s['massaPote']-s['gelatoPote']:.1f} g de cobertura" if s["fin"] and s["fin"][0]["adicional"] else ""))
    ws["A2"].font=PEQ

    # base de calculo: quanta CALDA existe no pote (ver nota no _dossie.js).
    # No Cafe a mescla desloca calda; no Alfajor o biscoito soma por cima.
    gel = s["caldaNoPote"]
    cab(ws, 4, ["Ingrediente","g por kg de calda","g no pote","R$/kg","R$ no pote","% do insumo"],
        [44,16,13,13,13,13])
    r=5; ini=r
    linhas = s["linhas"] + s["fin"]
    for l in linhas:
        ws.cell(row=r, column=1, value=l["nome"])
        ws.cell(row=r, column=2, value=round(l["gkg"],2)).number_format="0.0"
        # g no pote = g/kg x gelato/1000 ; a finalizacao usa a mesma base
        ws.cell(row=r, column=3, value=f"=B{r}*{gel:.4f}/1000").number_format="0.000"
        ws.cell(row=r, column=4,
                value=f"=IFERROR(VLOOKUP(A{r},{PRECO_TAB},2,FALSE),0)").number_format='"R$ "0.00'
        ws.cell(row=r, column=5, value=f"=C{r}*D{r}/1000").number_format='"R$ "0.000'
        r+=1
    fim=r-1
    ws.cell(row=r, column=1, value="INSUMO (sem perda)").font=NEG
    ws.cell(row=r, column=3, value=f"=SUM(C{ini}:C{fim})").number_format="0.00"
    ws.cell(row=r, column=5, value=f"=SUM(E{ini}:E{fim})").number_format='"R$ "0.000'
    for col in range(1,7): ws.cell(row=r, column=col).fill=FTOT
    ins_sem=r
    for rr in range(ini, fim+1):
        ws.cell(row=rr, column=6, value=f"=IFERROR(E{rr}/$E${ins_sem},0)").number_format="0.0%"
    r+=2

    ws.cell(row=r, column=1, value="FECHAMENTO").font=NEG; r+=1
    cab(ws, r, ["Linha","Valor","Como é calculado"], [34,16,52]); r+=1
    blocos=[
     ("Insumo com perda", f"=E{ins_sem}/(1-{PERDA_REF}/100)", "insumo ÷ (1 − perda de processo)"),
     ("Embalagem", f"={EMB_REF}", "pote + colher + lacre, da aba PARÂMETROS"),
     ("CMV TOTAL", None, "insumo com perda + embalagem"),
     ("Preço de venda", f"=VENDA!C{venda_lin[s['nome']]}", "da aba VENDA"),
     ("Imposto", None, "preço × alíquota do Simples"),
     ("Contribuição", None, "preço − CMV − imposto"),
     ("CMV % do preço", None, "CMV ÷ preço"),
     ("Margem de contribuição", None, "contribuição ÷ preço"),
    ]
    base=r
    for i,(n,fx,como) in enumerate(blocos):
        ws.cell(row=r, column=1, value=n)
        if n=="CMV TOTAL":            fx=f"=B{base}+B{base+1}"
        elif n=="Imposto":            fx=f"=B{base+3}*{IMP_REF}/100"
        elif n=="Contribuição":       fx=f"=B{base+3}-B{base+2}-B{base+4}"
        elif n=="CMV % do preço":     fx=f"=IFERROR(B{base+2}/B{base+3},0)"
        elif n=="Margem de contribuição": fx=f"=IFERROR(B{base+5}/B{base+3},0)"
        c=ws.cell(row=r, column=2, value=fx)
        c.number_format="0.0%" if "%" in n or "Margem" in n else '"R$ "0.00'
        ws.cell(row=r, column=3, value=como).font=PEQ
        if n in ("CMV TOTAL","Margem de contribuição"):
            for col in range(1,4): ws.cell(row=r, column=col).fill=FOK
            ws.cell(row=r, column=1).font=NEG; c.font=NEG
        r+=1
    aba_total[s["nome"]]=("CMV "+nome[:24], base+2, base+5, base+7, base+3)

# ═════════ 5. RESUMO
ws=wb.create_sheet("RESUMO", 3)
ws["A1"]="CMV e margem — os 7 produtos"; ws["A1"].font=TIT
ws["A2"]="Tudo puxado das abas de produto. Mude um preço na aba PREÇOS e esta tabela acompanha."
ws["A2"].font=PEQ
cab(ws, 4, ["Produto","Linha","Pote (g)","CMV","Venda","CMV %","Imposto","Contribuição","Margem"],
    [26,12,11,13,13,11,12,15,11])
r=5
for s in SABS:
    aba, lcmv, lctr, lmrg, lpv = aba_total[s["nome"]]
    q=f"'{aba}'!"
    ws.cell(row=r, column=1, value=s["nome"])
    ws.cell(row=r, column=2, value=s["linha"])
    ws.cell(row=r, column=3, value=round(s["massaPote"],0))
    ws.cell(row=r, column=4, value=f"={q}B{lcmv}").number_format='"R$ "0.00'
    ws.cell(row=r, column=5, value=f"={q}B{lpv}").number_format='"R$ "0.00'
    ws.cell(row=r, column=6, value=f"=IFERROR(D{r}/E{r},0)").number_format="0.0%"
    ws.cell(row=r, column=7, value=f"=E{r}*{IMP_REF}/100").number_format='"R$ "0.00'
    ws.cell(row=r, column=8, value=f"={q}B{lctr}").number_format='"R$ "0.00'
    ws.cell(row=r, column=9, value=f"=IFERROR(H{r}/E{r},0)").number_format="0.0%"
    for col in range(1,10): ws.cell(row=r, column=col).border=BF
    r+=1
ws.cell(row=r, column=1, value="MÉDIA DA LINHA").font=NEG
for col,fx in [(4,f"=AVERAGE(D5:D{r-1})"),(5,f"=AVERAGE(E5:E{r-1})"),
               (6,f"=AVERAGE(F5:F{r-1})"),(8,f"=AVERAGE(H5:H{r-1})"),
               (9,f"=AVERAGE(I5:I{r-1})")]:
    c=ws.cell(row=r, column=col, value=fx); c.font=NEG
    c.number_format="0.0%" if col in (6,9) else '"R$ "0.00'
for col in range(1,10): ws.cell(row=r, column=col).fill=FTOT
r+=2

ws.cell(row=r, column=1, value="SIMULAÇÃO DE PRODUÇÃO").font=NEG; r+=1
cab(ws, r, ["Produto","Potes a produzir","Calda (g)","Custo de insumo","Custo total","Receita","Contribuição"],
    [26,16,14,16,14,14,15]); r+=1
sim_ini=r
for i,s in enumerate(SABS):
    aba, lcmv, lctr, lmrg, lpv = aba_total[s["nome"]]
    q=f"'{aba}'!"
    lr = 5+i
    ws.cell(row=r, column=1, value=s["nome"])
    c=ws.cell(row=r, column=2, value=0); c.fill=FEDIT; c.border=BOX
    ws.cell(row=r, column=3,
        value=f"=B{r}*{s['caldaNoPote']:.4f}/(1-{PERDA_REF}/100)").number_format="#,##0"
    ws.cell(row=r, column=4, value=f"=B{r}*({q}B{lcmv}-{EMB_REF})").number_format='"R$ "#,##0.00'
    ws.cell(row=r, column=5, value=f"=B{r}*{q}B{lcmv}").number_format='"R$ "#,##0.00'
    ws.cell(row=r, column=6, value=f"=B{r}*E{lr}").number_format='"R$ "#,##0.00'
    ws.cell(row=r, column=7, value=f"=B{r}*{q}B{lctr}").number_format='"R$ "#,##0.00'
    for col in range(1,8): ws.cell(row=r, column=col).border=BF
    r+=1
ws.cell(row=r, column=1, value="TOTAL").font=NEG
for col in (2,3,4,5,6,7):
    c=ws.cell(row=r, column=col, value=f"=SUM({get_column_letter(col)}{sim_ini}:{get_column_letter(col)}{r-1})")
    c.font=NEG
    c.number_format="#,##0" if col in (2,3) else '"R$ "#,##0.00'
for col in range(1,8): ws.cell(row=r, column=col).fill=FTOT
r+=2
ws.cell(row=r, column=1, value="Preencha a coluna amarela com quantos potes de cada sabor e a planilha "
        "devolve insumo, custo total, receita e contribuição do lote.").font=PEQ
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)

wb.save("O Beco Gelato - CMV por Produto.xlsx")
print("CMV gerado — abas:", ", ".join(wb.sheetnames))
